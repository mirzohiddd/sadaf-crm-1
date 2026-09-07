"""Excel export — openpyxl."""
from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import crm

HEADER_FILL = PatternFill("solid", fgColor="0F2440")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)

# Hisobot turi -> (sarlavha, ustunlar [(kalit, nom)])
REPORTS: dict[str, dict[str, Any]] = {
    "leads": {
        "title": "Ledlar",
        "columns": [
            ("id", "ID"), ("name", "Led"), ("phone", "Telefon"), ("source", "Manba"),
            ("tour", "Tur"), ("people", "Odam"), ("amount", "Summa (USD)"),
            ("manager", "Mas'ul"), ("city", "Shahar"), ("date", "Sana"), ("stage", "Bosqich"),
        ],
    },
    "clients": {
        "title": "Mijozlar",
        "columns": [
            ("id", "ID"), ("name", "Mijoz"), ("phone", "Telefon"), ("country", "Mamlakat"),
            ("source", "Manba"), ("lastDate", "Oxirgi sayohat"), ("lastTour", "Oxirgi tur"),
            ("trips", "Sayohatlar"), ("total", "Jami xarajat"), ("manager", "Mas'ul"),
            ("status", "Status"),
        ],
    },
    "tours": {
        "title": "Turlar",
        "columns": [
            ("id", "ID"), ("name", "Tur nomi"), ("country", "Yo'nalish"), ("days", "Kun"),
            ("price", "Narx (USD)"), ("seats", "Bo'sh joy"), ("status", "Holat"),
        ],
    },
    "sales": {
        "title": "Savdolar",
        "columns": [
            ("id", "ID"), ("client", "Mijoz"), ("tour", "Tur"), ("amount", "Summa (USD)"),
            ("manager", "Menejer"), ("source", "Manba"), ("date", "Sana"),
        ],
    },
    "employees": {
        "title": "Hodimlar",
        "columns": [
            ("id", "ID"), ("name", "Hodim"), ("position", "Lavozim"), ("phone", "Telefon"),
            ("shift", "Ish vaqti"), ("hireDate", "Ishga kirgan"), ("deals", "Bitimlar"),
            ("status", "Holat"), ("login", "CRM login"), ("role", "CRM roli"),
        ],
    },
    "tasks": {
        "title": "Vazifalar",
        "columns": [
            ("id", "ID"), ("title", "Vazifa"), ("from", "Kimdan"), ("to", "Kimga"),
            ("priority", "Muhimlik"), ("due", "Muddat"), ("time", "Vaqt"), ("done", "Bajarildi"),
        ],
    },
    "attendance": {
        "title": "Davomat",
        "columns": [
            ("id", "ID"), ("userName", "Hodim"), ("date", "Sana"), ("checkIn", "Kirish"),
            ("checkOut", "Chiqish"), ("hours", "Soat"),
        ],
    },
}


def available_reports() -> list[dict[str, str]]:
    return [{"key": k, "title": v["title"]} for k, v in REPORTS.items()]


def rows_for(kind: str, user: dict[str, Any]) -> list[dict[str, Any]]:
    """Hisobot uchun ma'lumot — foydalanuvchi doirasida."""
    from .. import storage
    from ..deps import scope_rows

    if kind == "leads":
        return crm.leads_of(user)
    if kind == "clients":
        return crm.clients_of(user)
    if kind == "tours":
        return crm.tours_of(user)
    if kind == "sales":
        return crm.sales_of(user)
    if kind == "employees":
        return crm.employees_of(user)
    if kind == "tasks":
        return crm.tasks_of(user)
    if kind == "attendance":
        from ..deps import is_global, visible_ids

        rows = storage.read("attendance")
        if is_global(user):
            return rows
        ids = visible_ids(user) or set()
        return [r for r in rows if int(r.get("userId", 0)) in ids]
    return []


def _cell(value: Any) -> Any:
    if isinstance(value, bool):
        return "Ha" if value else "Yo'q"
    if isinstance(value, (list, dict)):
        return str(value)
    return value


def build_sheet_data(kind: str, user: dict[str, Any]) -> tuple[list[str], list[list[Any]]]:
    """Excel eksport uchun umumiy: sarlavhalar + qatorlar."""
    spec = REPORTS.get(kind)
    if not spec:
        raise ValueError(f"Noma'lum hisobot turi: {kind}")
    columns = spec["columns"]
    headers = [label for _, label in columns]
    rows = [[_cell(r.get(key, "")) for key, _ in columns] for r in rows_for(kind, user)]
    return headers, rows


def build_workbook(kinds: list[str], user: dict[str, Any]) -> bytes:
    """Bir yoki bir nechta hisobotni bitta .xlsx faylga yig'adi."""
    wb = Workbook()
    wb.remove(wb.active)

    for kind in kinds:
        if kind not in REPORTS:
            continue
        headers, rows = build_sheet_data(kind, user)
        ws = wb.create_sheet(title=REPORTS[kind]["title"][:31])

        ws.append(headers)
        for cell in ws[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.freeze_panes = "A2"

        for row in rows:
            ws.append(row)

        # Ustun kengligini mazmunga moslash
        for idx, header in enumerate(headers, start=1):
            longest = max(
                [len(str(header))] + [len(str(r[idx - 1])) for r in rows[:200]] or [10]
            )
            ws.column_dimensions[get_column_letter(idx)].width = min(38, max(10, longest + 3))

        if rows:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"

    if not wb.sheetnames:
        ws = wb.create_sheet(title="Bo'sh")
        ws.append(["Ma'lumot topilmadi"])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
